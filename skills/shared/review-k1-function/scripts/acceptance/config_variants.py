"""Modify only local copies of frozen K1 workbooks; row 3 names and row 9 data."""
from copy import copy
from decimal import Decimal, ROUND_FLOOR
from pathlib import Path
import json
from openpyxl import load_workbook
from .core import require, file_hash, reject_formal_write, owned_path, atomic_json


def aggregate_floor(values): return int(sum((Decimal(str(v)) for v in values), Decimal(0)).to_integral_value(rounding=ROUND_FLOOR))
def divide_floor(value, divisor):
    require(type(value) in [str, int] and type(divisor) is int and divisor > 0, 'INTEGER_REQUIRED')
    return int(value) // divisor


def generate(recipe, sources, target, output):
    require(recipe.get('schemaVersion') == 1, 'SCHEMA_VERSION')
    output = Path(output).resolve(); reject_formal_write(output); require(not output.exists(), 'BUNDLE_EXISTS')
    loaded = {}; changes = []
    for key, source in sources.items():
        require(file_hash(source['path']) == source['sha256'], 'SOURCE_HASH_MISMATCH')
        loaded[key] = load_workbook(source['path'])
    for change in recipe['changes']:
        w = loaded[change['workbook']]; s = w[change['sheet']]
        fields = {c.value: c.column for c in s[3] if c.value is not None}
        require(change['keyField'] in fields and change['field'] in fields, 'FIELD_NOT_FOUND')
        keycol = fields[change['keyField']]
        rows = [r for r in range(9, s.max_row + 1) if str(s.cell(r, keycol).value) == str(change['key'])]
        require(len(rows) == 1, 'KEY_AMBIGUOUS_OR_MISSING')
        row = rows[0]
        if 'newKey' in change:
            require(not any(str(s.cell(r, keycol).value) == str(change['newKey']) for r in range(9, s.max_row + 1)), 'ID_COLLISION')
            newrow = s.max_row + 1
            for c in range(1, s.max_column + 1): s.cell(newrow, c).value = s.cell(row, c).value; s.cell(newrow, c)._style = copy(s.cell(row, c)._style)
            row = newrow; s.cell(row, keycol).value = change['newKey']
        col = fields[change['field']]; old = s.cell(row, col).value; value = change['value']
        if change.get('requireDifferent'): require(str(value) != str(old) and str(value) != str(change.get('currentlyLoaded')), 'VARIANT_NOT_DIFFERENT')
        if type(value) is int and abs(value) > 2**53: value = str(value)  # Excel doubles must not round large integer scenarios.
        s.cell(row, col).value = value
        changes.append({**change, 'before': old, 'after': value, 'row': row})
    output.mkdir(parents=True, exist_ok=False); files = []
    for key, w in loaded.items():
        source = sources[key]; dest = owned_path(output, key + '.xlsx'); w.save(dest); w.close()
        require(file_hash(source['path']) == source['sha256'], 'SOURCE_CHANGED_DURING_GENERATION')
        files.append({'workbook': key, 'path': str(dest), 'sha256': file_hash(dest), 'sourceSha256': source['sha256']})
    summary = {'schemaVersion': 1, 'files': files, 'changes': changes, 'target': target, 'liveValidation': 'not_deployed'}
    atomic_json(output / 'summary.json', summary)
    (output / '表结构与差异.md').write_text('# 本地测试副本\n\n保留正式表前八行和现有主键名；未生成新的展示文案，不需要新增语言表。\n\n```json\n' + json.dumps(changes, ensure_ascii=False, indent=2) + '\n```\n', encoding='utf-8')
    return summary
