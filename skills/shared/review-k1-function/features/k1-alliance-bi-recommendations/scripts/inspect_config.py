import json
import sys
from pathlib import Path

from openpyxl import load_workbook

root = Path(sys.argv[1])
queries = {
    "ConstConfig-materialized-target.xlsx": {str(v) for v in range(10135700, 10135752)},
    "NewMail-materialized-target.xlsx": {"10950114", "20000012", "20000015", "17112009", "17112010"},
    "Switch-materialized-target.xlsx": {str(v) for v in range(336, 344)},
}
result = {}
for name, ids in queries.items():
    path = root / name
    wb = load_workbook(path, read_only=True, data_only=True)
    workbook = {"sha256": None, "sheets": {}}
    for ws in wb.worksheets:
        headers = [ws.cell(3, col).value for col in range(1, ws.max_column + 1)]
        matches = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=9, values_only=True), start=9):
            if row and str(row[0]).strip() in ids:
                matches.append({
                    "row": row_idx,
                    "values": {
                        str(headers[i] or f"col_{i+1}"): value
                        for i, value in enumerate(row)
                        if value is not None and str(value) != ""
                    },
                })
        if matches:
            workbook["sheets"][ws.title] = {
                "dimensions": [ws.max_row, ws.max_column],
                "header_rows": [[ws.cell(r, c).value for c in range(1, ws.max_column + 1)] for r in range(1, 9)],
                "matches": matches,
            }
    result[name] = workbook
print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
