import json, hashlib
from pathlib import Path
from openpyxl import load_workbook
base=Path('__CONFIG_SNAPSHOT_DIRECTORY__')
out=[]
for name in ['ColdServer.xlsx','ConstConfig.xlsx','ActivityOnline.xlsx']:
 p=base/name
 wb=load_workbook(p,read_only=True,data_only=False)
 rec={'file':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'sheets':[]}
 for ws in wb:
  rows=list(ws.iter_rows(values_only=True)); hits=[]
  for i,r in enumerate(rows[8:],9):
   s=' '.join(str(v) for v in r if v is not None)
   if (name=='ColdServer.xlsx' and any(str(v) in ['__SOURCE_SERVER_ID__','__TARGET_SERVER_B_ID__','__TARGET_SERVER_C_ID__'] for v in r)) or (name=='ConstConfig.xlsx' and any(k in s for k in ['10135787','10135788','10135790','10135791','冷库','冷冻'])) or (name=='ActivityOnline.xlsx' and any(str(v)=='__ACTIVITY_TEMPLATE_ID__' for v in r[:4])):
    hits.append({'row':i,'values':r})
  rec['sheets'].append({'sheet':ws.title,'rows':ws.max_row,'headers':rows[:8],'matches':hits})
 out.append(rec)
 wb.close()
dest=Path('__OUTPUT_FILE__')
dest.write_text(json.dumps(out,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
print(json.dumps([{'file':x['file'],'sheets':[{'sheet':s['sheet'],'matches':s['matches']} for s in x['sheets']]} for x in out],ensure_ascii=False,default=str))
